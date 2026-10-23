"""Numerical and end-to-end checks, independent of MATLAB and the NN package."""
import csv
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import numpy as np
from numpy.testing import assert_allclose
from openpyxl import Workbook, load_workbook

from gidind_reference.config import Settings
from gidind_reference.geometry import (assign, cartesian_batches, direct_metric, hkl_grid,
    metric_design, metric_to_cell, observed3, predict, primitive_planes, reciprocal_metric,
    reduce_basis, first_order, matlab_round)
from gidind_reference.inputs import read_peaks, seed_coordinates
from gidind_reference.runner import index_file
from gidind_reference.specular import abg_from_xyz, stage1_matrices
from gidind_reference.without_specular import refit, seed_systems
from gidind_reference.state import State, Candidate
from gidind_reference.finalize import finalize

ROOT = Path(__file__).resolve().parent/'fixtures'


def synthetic():
    cell = np.array([5.,7.,9.,np.pi/2,np.pi/2,np.pi/2])
    metric = reciprocal_metric(cell)
    h = np.array([[0,1,1],[1,0,1],[1,1,1],[-1,1,1],[0,1,2],[1,0,2],
                  [1,1,2],[-1,1,2],[1,2,1],[2,1,1],[2,0,1],[1,0,3]])
    plane = np.array([0,0,1])
    return cell, metric, plane, h, predict(metric,h,plane=plane)[:,:2]


def write_csv(path, coords, header=True):
    with Path(path).open('w',newline='') as f:
        w=csv.writer(f)
        if header: w.writerow(['q_xy','q_z'])
        w.writerows(coords)


class InputTests(unittest.TestCase):
    def test_csv_excel_order_and_filter_provenance(self):
        _,_,_,_,coords=synthetic()
        coords=np.vstack([[0,0],[0,1],coords])
        with tempfile.TemporaryDirectory() as tmp:
            csvpath=Path(tmp)/'a.csv'
            with csvpath.open('w',newline='') as f:
                w=csv.writer(f); w.writerow(['id','q_z','q_xy'])
                for i,(x,z) in enumerate(coords): w.writerow([i,z,x])
            xlsx=Path(tmp)/'b.xlsx'; wb=Workbook(); wb.active.append(['q_z','q_xy'])
            for x,z in coords: wb.active.append([float(z),float(x)])
            wb.save(xlsx)
            a=read_peaks(csvpath,'without-specular',Settings())
            b=read_peaks(xlsx,'without-specular',Settings())
            assert_allclose(a.coordinates,b.coordinates)
            self.assertEqual(a.provenance['removed_rows'],[
                {'source_row':2,'reason':'zero_placeholder'}, {'source_row':3,'reason':'specular_filter'}])
            self.assertEqual(a.source_rows[0],4)
            spec=read_peaks(csvpath,'specular',Settings())
            self.assertEqual(spec.provenance['selected_specular_q'],1)
            self.assertEqual(spec.provenance['selected_specular_source_row'],3)

    def test_headerless_first_row_preserved_and_invalid_rejected(self):
        _,_,_,_,coords=synthetic()
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'a.csv'; write_csv(p,coords,False)
            self.assertEqual(len(read_peaks(p,'without-specular',Settings()).coordinates),len(coords))
            for bad in ([np.nan,1],[-1,1],[1,np.inf],['',1]):
                write_csv(p,np.vstack([coords,np.array(bad,dtype=object)]))
                with self.assertRaises(ValueError): read_peaks(p,'without-specular',Settings())
            write_csv(p,coords)
            with self.assertRaisesRegex(ValueError,'requires a positive'):
                read_peaks(p,'specular',Settings())

    def test_configuration_rejects_invalid_controls(self):
        for kwargs in ({'batch_size':0},{'plane_limit':-1},{'time_limit_seconds':float('nan')},
                       {'planes':[[0,0,0]]},{'rmse_cutoffs':[0,1,1]},
                       {'without_specular_seed_hkl':[1,1,-1]}):
            with self.assertRaises(ValueError): Settings(**kwargs).validate()


class GeometryTests(unittest.TestCase):
    def test_reciprocal_metric_against_cartesian_vectors(self):
        cell=np.array([5.1,7.2,9.3,*np.deg2rad([83,101,107])])
        a,b,c,al,be,ga=cell
        direct=np.array([[a,b*np.cos(ga),c*np.cos(be)],
                         [0,b*np.sin(ga),c*(np.cos(al)-np.cos(be)*np.cos(ga))/np.sin(ga)],
                         [0,0,0.]])
        direct[2,2]=np.sqrt(c*c-direct[0,2]**2-direct[1,2]**2)
        reciprocal=2*np.pi*np.linalg.inv(direct).T
        metric=reciprocal_metric(cell)
        assert_allclose(metric,reciprocal.T@reciprocal,atol=1e-13)
        assert_allclose(metric_to_cell(metric),cell,atol=1e-12)
        h=np.array([[1,-1,2],[0,0,1],[1,1,0]])
        plane=np.array([1,0,2]); normal=reciprocal@plane; normal/=np.linalg.norm(normal)
        vectors=h@reciprocal.T
        z=vectors@normal; xyz=np.linalg.norm(vectors,axis=1)
        expected=np.column_stack([np.sqrt(np.maximum(xyz*xyz-z*z,0)),z,xyz])
        assert_allclose(predict(metric,h,plane=plane),expected,atol=2e-8)

    def test_enumeration_and_bounded_integer_planes(self):
        a=np.vstack([b for _,b in cartesian_batches([-1,0,1],3,5)])
        b=np.vstack([b for _,b in cartesian_batches([-1,0,1],3,27)])
        assert_allclose(a,b)
        assert_allclose(a[:4],[[-1,-1,-1],[0,-1,-1],[1,-1,-1],[-1,0,-1]])
        planes=primitive_planes(2)
        self.assertEqual(len(planes),49)
        self.assertTrue(all(np.gcd.reduce(abs(p))==1 for p in planes))
        assert_allclose(primitive_planes(explicit=[[0,0,2],[0,0,-1]]),[[0,0,1]])
        assert_allclose(matlab_round([-.5,.5,-1.5,1.5]),[-1,1,-2,2])

    def test_specular_seed_equations_recover_exact_abgamma(self):
        cell,metric,plane,h,coords=synthetic()
        qspec=float(np.sqrt(plane@metric@plane))
        matrices=stage1_matrices(h[:3,:2].ravel()[None,:],0,0,coords[:3],qspec)
        xyz=np.linalg.solve(matrices[0],np.ones(3))
        found=abg_from_xyz(xyz[None,:],0,0,qspec,Settings())
        assert_allclose(found,cell[[0,1,5]][None,:],atol=1e-10)

    def test_no_specular_seed_equations_and_refit(self):
        cell,metric,plane,h,coords=synthetic()
        cell[3:]=np.deg2rad([93,98,104]); metric=reciprocal_metric(cell)
        coords=predict(metric,h,plane=plane)[:,:2]
        order=np.argsort(np.linalg.norm(coords,axis=1),kind='stable')
        seed=coords[order][:8]; seedh=h[order][:8]
        # Exercise the actual AA4/fifth/sixth selection on the known assignment.
        _,h6,selected,_=seed_systems(seedh[:4][None,:,:],seed,hkl_grid((3,3,3),sorted_norm=True))
        self.assertEqual(len(h6),1)
        design=metric_design(h6[0]); self.assertEqual(np.linalg.matrix_rank(design),6)
        from gidind_reference.geometry import six_to_metric
        recovered=six_to_metric(np.linalg.solve(design,np.sum(seed[selected[0]]**2,axis=1)))
        assert_allclose(recovered,metric,atol=1e-10)
        projection=metric@plane/np.sqrt(plane@metric@plane)
        candidate,why=refit(metric,projection,coords,hkl_grid((3,3,3)),.1)
        self.assertIsNone(why)
        assert_allclose(candidate.metric,metric,atol=1e-10)

    def test_basis_reduction_transforms_reflections_and_plane_together(self):
        cell,metric,plane,h,coords=synthetic()
        # Nontrivial integer shear of a known cell before running native reduction.
        transform=np.array([[1,1,0],[0,1,0],[0,0,1]])
        inv=np.linalg.inv(transform)
        sheared=inv.T@metric@inv
        hh=h@transform.T; pp=transform@plane
        reduced,newh,newplane,t=reduce_basis(sheared,hh,plane=pp)
        assert_allclose(newh,hh@t.T,atol=1e-10)
        assert_allclose(newplane,t@pp,atol=1e-10)
        assert_allclose(predict(reduced,newh,plane=newplane),predict(sheared,hh,plane=pp),atol=1e-8)
        assert_allclose(np.linalg.det(reduced),np.linalg.det(sheared),rtol=1e-10)

    def test_native_first_order_reduces_small_cell_perturbation(self):
        cell,metric,plane,h,coords=synthetic()
        perturbed=cell.copy(); perturbed[:3]*=[1.001,.999,1.002]
        before=predict(reciprocal_metric(perturbed),h,plane=plane)-observed3(coords)
        after,applied=first_order(perturbed,h,coords,plane=plane)
        # Native determinant gate is allowed to decline a correction.
        if applied:
            self.assertLess(np.linalg.norm(predict(reciprocal_metric(after),h,plane=plane)-observed3(coords)),np.linalg.norm(before))
        else:
            assert_allclose(after,perturbed)

    def test_bundled_pq_export_reconstructs_reported_qxyz_error(self):
        wb=load_workbook(ROOT/'Exported_example_PQ_75.xlsx',read_only=True,data_only=True)
        rows=list(wb.active.values); wb.close()
        cell=np.r_[[float(x) for x in rows[1][:3]],np.deg2rad([float(x) for x in rows[3][:3]])]
        plane=np.array([float(x) for x in rows[5][:3]])
        points=np.array([[float(x) for x in row] for row in rows[11:]])
        errors=np.sqrt(np.mean((predict(reciprocal_metric(cell),points[:,2:],plane=plane)-observed3(points[:,:2]))**2,axis=0))
        self.assertLess(abs(errors[2]-float(rows[8][0])),.0001)
        self.assertLess(errors.max(),.003)


class IntegrationTests(unittest.TestCase):
    def test_benchmark_reuses_prepared_inputs_and_remaps_source_rows(self):
        from gidind_reference.batch import benchmark
        cell=np.r_[[5.,7.,9.],np.deg2rad([93,98,104])]
        metric=reciprocal_metric(cell)
        pp=predict(metric,hkl_grid((1,1,1)),plane=[0,0,1])
        coords=pp[(pp[:,1]>.05)&(pp[:,0]>.05),:2]
        with tempfile.TemporaryDirectory() as tmp:
            inputs=Path(tmp)/'inputs';inputs.mkdir();cases=[]
            for case in ('pq','mapbbr3','naib_s11','cu_bod_mof'):
                original=inputs/f'{case}.csv';filtered=inputs/f'{case}_without_specular.csv'
                write_csv(original,np.vstack([[0,0],coords]))
                write_csv(filtered,coords)
                digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
                cases.append({'id':case,'filename':original.name,'without_specular':{
                    'filename':filtered.name,'sha256':digest(filtered),'source_sha256':digest(original)}})
            (inputs/'manifest.json').write_text(json.dumps({'cases':cases}))
            settings=Settings(algorithm='upstream',without_specular_seed_hkl=[1,1,1],planes=[[0,0,1]],time_limit_seconds=60,
                              rmse_cutoffs=[1e-6]*3)
            summary=benchmark(inputs,Path(tmp)/'out',settings)
            self.assertEqual(summary['unique_computations'],1)
            self.assertEqual(len(summary['jobs']),8)
            self.assertEqual(sum(j['reused'] for j in summary['jobs']),7)
            alias=json.loads((Path(tmp)/'out/pq_without_specular/results.json').read_text())
            self.assertEqual(alias['solutions'][0]['assignments'][0]['source_row'],2)
            self.assertFalse(alias['deduplication']['independent_timing_measurement'])
            self.assertTrue((Path(tmp)/'out/pq_without_specular/port_source_snapshot.zip').is_file())

    def test_interrupted_search_preserves_native_candidate(self):
        _,metric,plane,h,coords=synthetic()
        def stop(peaks, seed, settings, state):
            state.add_pending(Candidate(metric,h,None,metric@plane/np.sqrt(plane@metric@plane),
                                        {'seed_linear_sumerr':0.}))
            raise KeyboardInterrupt()
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'a.csv';write_csv(path,coords)
            with patch('gidind_reference.without_specular.search',side_effect=stop):
                result=index_file(path,'without-specular',Path(tmp)/'out',Settings())
            self.assertEqual(result['status'],'interrupted')
            self.assertEqual(result['pending_candidates'],1)
            with np.load(Path(tmp)/'out/intermediate_candidates.npz',allow_pickle=False) as arrays:
                assert_allclose(arrays['reciprocal_metrics'][0],metric)
                assert_allclose(arrays['hkl'][0],h)

    def test_full_no_specular_search_recovers_synthetic_triclinic_lattice(self):
        cell=np.r_[[5.,7.,9.],np.deg2rad([93,98,104])]
        metric=reciprocal_metric(cell);plane=np.array([0,0,1])
        predicted=predict(metric,hkl_grid((1,1,1)),plane=plane)
        coords=predicted[(predicted[:,1]>.05)&(predicted[:,0]>.05),:2]
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'a.csv';write_csv(path,coords)
            settings=Settings(algorithm='upstream',without_specular_seed_hkl=[1,1,1],planes=[[0,0,1]],
                              time_limit_seconds=60,rmse_cutoffs=[1e-6]*3)
            result=index_file(path,'without-specular',Path(tmp)/'out',settings)
            self.assertEqual(result['status'],'completed',result.get('error',result['counts']))
            self.assertTrue(result['solutions'])
            actual=result['solutions'][0]['cell']
            assert_allclose([actual[k] for k in ('a','b','c')],cell[:3],atol=1e-7)
            assert_allclose([actual[k] for k in ('alpha','beta','gamma')],np.rad2deg(cell[3:]),atol=1e-7)

    def test_seed_batch_size_does_not_change_no_specular_systems(self):
        cell,metric,plane,h,coords=synthetic()
        cell[3:]=np.deg2rad([93,98,104]);metric=reciprocal_metric(cell)
        coords=predict(metric,h,plane=plane)[:,:2]
        order=np.argsort(np.linalg.norm(coords,axis=1),kind='stable')[:8]
        seed=coords[order];hk=hkl_grid((1,1,1),sorted_norm=True)
        # Use explicit valid and rank-deficient H4 systems in mixed batches.
        h4=np.array([h[order[:4]],h[order[1:5]],np.zeros((4,3)),h[order[2:6]]])
        _,allh,allsel,allq=seed_systems(h4,seed,hk)
        pieces=[seed_systems(h4[i:i+1],seed,hk) for i in range(len(h4))]
        assert_allclose(allh,np.concatenate([p[1] for p in pieces]))
        assert_allclose(allsel,np.concatenate([p[2] for p in pieces]))
        assert_allclose(allq,np.concatenate([p[3] for p in pieces]))

    def test_specular_search_recovers_synthetic_lattice(self):
        cell,metric,plane,h,coords=synthetic()
        coords=np.vstack([[0,np.sqrt(plane@metric@plane)],coords])
        settings=Settings(specular_seed_peaks=4,specular_seed_hk=1,specular_index_hk=2,specular_l=3,specular_uv=1,
                          specular_w=1,redundancy_hk=2,planes=[[0,0,1]],batch_size=729,
                          time_limit_seconds=30,rmse_cutoffs=[1e-6]*3,
                          max_intermediate_candidates=2000)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'a.csv';write_csv(path,coords)
            result=index_file(path,'specular',Path(tmp)/'out',settings)
            self.assertNotEqual(result['status'],'failure',result.get('error'))
            self.assertTrue(result['solutions'],result['counts'])
            self.assertLess(result['solutions'][0]['rmse']['q_xyz'],1e-7)
            self.assertTrue(any(np.allclose([s['cell'][k] for k in ('a','b','c')],cell[:3],atol=1e-5)
                                for s in result['solutions']))
            self.assertIsNone(result['parameter_uncertainty'])
            self.assertTrue((Path(tmp)/'out/assignments/solution_00001.csv').is_file())
            from dataclasses import replace
            other=index_file(path,'specular',Path(tmp)/'other_batch',replace(settings,batch_size=137))
            self.assertEqual(other['status'],'completed')
            signature=lambda r: sorted((tuple(round(s['cell'][k],6) for k in ('a','b','c','alpha','beta','gamma')),
                                       tuple(s['contact_plane'])) for s in r['solutions'])
            self.assertEqual(signature(result),signature(other))
            with self.assertRaises(FileExistsError): index_file(path,'specular',Path(tmp)/'out',settings)

    def test_no_specular_finalization_and_exact_zero_window(self):
        _,metric,plane,h,coords=synthetic()
        settings=Settings(planes=[[0,0,1]],time_limit_seconds=10)
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'a.csv';write_csv(p,coords)
            peaks=read_peaks(p,'without-specular',settings)
            candidate,_=refit(metric,metric@plane/np.sqrt(plane@metric@plane),coords,hkl_grid((3,3,3)),.1)
            candidate.native_scores['seed_linear_sumerr']=0.
            state=State(settings);state.add_pending(candidate)
            finalize(peaks,'without-specular',settings,state)
            self.assertTrue(state.solutions)
            self.assertLess(state.solutions[0]['rmse']['q_xyz'],1e-7)

    def test_finalization_deduplicates_identical_indexed_solutions(self):
        _,metric,plane,h,coords=synthetic()
        settings=Settings(planes=[[0,0,1]],time_limit_seconds=10)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'a.csv';write_csv(path,coords)
            peaks=read_peaks(path,'without-specular',settings)
            candidate,_=refit(metric,metric@plane/np.sqrt(plane@metric@plane),coords,hkl_grid((3,3,3)),.1)
            state=State(settings)
            state.pending=[candidate,candidate]
            finalize(peaks,'without-specular',settings,state)
            self.assertEqual(len(state.solutions),1)
            self.assertEqual(state.counts['duplicate_final_solution'],1)

    def test_time_limit_checkpoint_and_input_failure_are_explicit(self):
        _,_,_,_,coords=synthetic()
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'a.csv';write_csv(p,coords)
            result=index_file(p,'without-specular',Path(tmp)/'limited',Settings(time_limit_seconds=.001))
            self.assertEqual(result['status'],'time_limited')
            self.assertFalse(result['search_complete'])
            self.assertTrue((Path(tmp)/'limited/checkpoint.json').is_file())
            failed=index_file(p,'specular',Path(tmp)/'failed',Settings())
            self.assertEqual(failed['status'],'failure')
            self.assertIn('requires a positive',failed['error'])


if __name__=='__main__':
    unittest.main()
